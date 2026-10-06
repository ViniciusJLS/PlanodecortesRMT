"""Fluxo Excel → configuração do lançamento → prévia → cálculo."""
import base64
import hashlib
import sqlite3
import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter
from . import engine, project as projects
from .reference import (FIELDS, sheets, read_sheet, guess_header, guess_mapping, parse_reference,
                        apply_reference, register_park, reference_template, optimize_scope, scope_key)
from .subparks import CIRCUITS, LEVELS
from .exporter import export, export_subpark
from .tramo import is_aux, aux_bounds, parse_aux
from .importer import norm
from zipfile import BadZipFile


@st.cache_data(show_spinner=False)
def sheet_data(content, sheet):
    return read_sheet(content, sheet)


def persist(project, candidate, message):
    projects.save(candidate)
    project.clear()
    project.update(candidate)
    st.session_state.edit_version += 1
    st.session_state.reference_message = message
    st.rerun()


def new_park(project):
    with st.expander('Cadastrar subparque', expanded=not project['trechos'] and not project.get('subparques')):
        with st.form(f"new_park_{project['id']}"):
            name = st.text_input('Nome do novo subparque', placeholder='Ex.: RSA-03')
            if st.form_submit_button('Cadastrar e selecionar'):
                try:
                    candidate, name = register_park(project, name)
                    projects.save(candidate)
                except (ValueError, OSError, sqlite3.Error) as exc:
                    st.error(str(exc))
                else:
                    project.clear()
                    project.update(candidate)
                    st.session_state[f"subpark_{project['id']}_park"] = name
                    st.rerun()


def render_reference(project, park):
    prefix = f"reference_{project['id']}_{park}"
    if message := st.session_state.pop('reference_message', None):
        st.success(message)
    with st.expander('Configurar lançamento e importar Excel deste subparque', expanded=True):
        st.caption('Estes seletores definem o lançamento importado. Os filtros da consulta apenas mudam a visualização.')
        a,b,c = st.columns(3)
        circuit = a.selectbox('Circuito do lançamento', CIRCUITS, key=prefix+'_circuit')
        level = b.selectbox('Nível do lançamento', LEVELS, key=prefix+'_level')
        typ = c.selectbox('Instalação do lançamento', ['AÉREO','SUBTERRÂNEO'], key=prefix+'_type')
        a,b = st.columns(2)
        route = a.text_input('Rota do lançamento', 'Principal', key=prefix+'_route')
        conductors = sorted({r['condutor'] for r in project['bobinas'] if r.get('tipo')==typ})
        conductor = b.selectbox('Condutor padrão', ['Ler da coluna do Excel']+conductors+['Informar outro condutor'], key=prefix+'_conductor_'+typ)
        if conductor == 'Informar outro condutor':
            conductor = st.text_input('Descrição do condutor', key=prefix+'_custom_conductor')
        elif conductor == 'Ler da coluna do Excel':
            conductor = ''
        a,b,c = st.columns(3)
        percent = a.number_input('Folga padrão [%]', min_value=0.0, max_value=100.0,
            value=2.0 if typ=='SUBTERRÂNEO' else 5.0, key=prefix+'_allowance_'+typ)
        reserve = b.number_input('Reserva padrão por fase e trecho [m]', min_value=0.0, value=0.0, key=prefix+'_reserve')
        cut = c.selectbox('Corte padrão no fim', ['PERMITIDO','PROIBIDO','OBRIGATORIO'], key=prefix+'_cut')
        st.caption('Os padrões são usados quando não há coluna associada. Confirme amarrações e travessias antes de otimizar.')
        st.download_button('Baixar modelo Excel de referência', reference_template(), 'Modelo_referencia_subparque.xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', key=prefix+'_template')
        uploaded = st.file_uploader(f'Excel de referência — {park}', type=['xlsx','xlsm'], key=prefix+'_upload')
        if uploaded is not None:
            content = uploaded.getvalue()
            if len(content)>10*1024*1024:
                st.error('Use um Excel com até 10 MB, contendo apenas as tabelas de referência necessárias.')
            else:
                try:
                    render_mapping(project, content, uploaded.name, prefix,
                        dict(parque=park,circuito=circuit,nivel=level,tipo=typ,rota=route.strip(),
                             condutor=conductor,folga=percent/100,reserva=reserve,corte_fim=cut))
                except (ValueError, KeyError, OSError, BadZipFile) as exc:
                    st.error(str(exc))
    references = project.get('referencias_subparques',{}).get(park,{})
    if references:
        with st.expander(f'Referências salvas de {park} ({len(references)})'):
            for scope, ref in references.items():
                st.write(f"**{scope}** — {ref['filename']} · aba {ref['sheet']} · {ref['registros']} registros de fase")
                content = project.get('arquivos_referencia',{}).get(ref['sha256'])
                if content:
                    st.download_button('Baixar referência original', base64.b64decode(content), ref['filename'],key=prefix+'_download_'+scope)


def render_mapping(project, content, filename, prefix, config):
    digest = hashlib.sha256(content).hexdigest()
    key = prefix+'_'+digest[:16]
    names = sheets(content)
    preferred = config['parque'] if config['parque'] in names else 'PLANO DE CORTE' if 'PLANO DE CORTE' in names else names[0]
    sheet = st.selectbox('Aba do Excel', names, index=names.index(preferred), key=key+'_sheet')
    config['sheet'] = sheet
    data = sheet_data(content, sheet)
    if not data:
        st.warning('A aba selecionada está vazia.')
        return
    key += '_'+sheet
    detected = any(is_aux(row) for row in data[:50])
    layout = st.selectbox('Formato da referência', ['Tabela com De e Para','AUX TRAMO: postes em I e distância em L'],
        index=1 if detected else 0, key=key+'_layout')
    if layout.startswith('AUX TRAMO'):
        render_aux(project,content,filename,key,config,data)
        return
    header = int(st.number_input('Linha do cabeçalho', min_value=1, max_value=len(data), value=guess_header(data), key=key+'_header'))
    if header == len(data):
        st.warning('Não há linhas depois do cabeçalho.')
        return
    key += '_'+str(header)+'_'+config['nivel']
    headers = data[header-1]
    guessed = guess_mapping(headers, config['nivel'])
    options = [None]+list(range(len(headers)))
    def label(index):
        return 'Não associar / usar padrão' if index is None else f'{get_column_letter(index+1)} · {headers[index] or "Sem título"}'
    mapping = {}
    st.write('**Associe as colunas da referência**')
    cols = st.columns(3)
    for i, field in enumerate(('de','para','linear','fase','condutor')):
        mapping[field] = cols[i%3].selectbox(FIELDS[field], options, index=options.index(guessed[field]), format_func=label, key=key+'_'+field)
    with st.expander('Demais colunas: circuito, tipo, folga, reservas e cortes'):
        cols = st.columns(3)
        for i, field in enumerate(f for f in FIELDS if f not in mapping):
            mapping[field] = cols[i%3].selectbox(FIELDS[field], options, index=options.index(guessed[field]), format_func=label, key=key+'_'+field)
        st.caption('Circuito, nível e tipo associados ao Excel filtram as linhas para a configuração escolhida. Sem associação, as linhas recebem a configuração acima. Mapeie a reserva total OU as sobras detalhadas.')
    a,b = st.columns(2)
    start = int(a.number_input('Primeira linha de dados', min_value=header+1, max_value=len(data), value=header+1, key=key+'_start'))
    end = int(b.number_input('Última linha de dados', min_value=header+1, max_value=len(data), value=len(data), key=key+'_end'))
    with st.expander('Visualizar linhas originais do intervalo'):
        preview = pd.DataFrame(data[start-1:min(end,start+298)],
            columns=[f'{get_column_letter(i+1)} · {h or "Sem título"}' for i,h in enumerate(headers)])
        preview.index = range(start,start+len(preview))
        st.dataframe(preview.fillna('').astype(str), width='stretch')
        st.caption('O índice à esquerda é o número da linha no Excel. Ajuste o intervalo para excluir totais, notas e blocos que não pertencem ao lançamento.')
    mode = st.radio('Organização das fases no Excel', ['Uma linha por trecho: gerar A, B e C','Uma linha por fase'],
        index=1 if mapping['fase'] is not None else 0, key=key+'_phase_mode')
    phase_mode = 'column' if mode=='Uma linha por fase' else 'expand'
    percentage = st.checkbox('A coluna de folga contém 5 para representar 5% (desmarcado: 0,05)', key=key+'_percentage')
    st.caption('Associe a distância linear original, não a distância de projeto arredondada. Fórmulas precisam estar calculadas e salvas no Excel.')
    if mapping.get('linear') is not None and 'PROJETO' in norm(headers[mapping['linear']]):
        st.error('A coluna escolhida contém distância de projeto. Associe a distância linear original.')
        return
    try:
        rows, skipped = parse_reference(data, mapping, config, start, end, phase_mode, percentage)
    except ValueError as exc:
        st.error(str(exc))
        return
    preview_and_import(project, content, filename, key, config, rows, mapping, start, end, phase_mode, percentage, skipped)


def render_aux(project,content,filename,key,config,data):
    _,first,last=aux_bounds(data)
    st.info('Coluna I: poste. Coluna L: distância entre o poste da linha anterior (De) e o da linha atual (Para). Distância zero inicia outro tramo; linhas vazias e marcadores P. não viram trechos.')
    a,b=st.columns(2)
    start=int(a.number_input('Primeira linha de postes',min_value=1,max_value=len(data),value=first,key=key+'_aux_start'))
    end=int(b.number_input('Última linha de postes',min_value=1,max_value=len(data),value=last,key=key+'_aux_end'))
    with st.expander('Visualizar postes I e distâncias L'):
        preview=[{'Linha Excel':i,'Poste (I)':str(data[i-1][8] or ''),'Distância tramo (L)':str(data[i-1][11])} for i in range(start,min(end,start+299)+1)]
        st.dataframe(pd.DataFrame(preview),hide_index=True,width='stretch')
    try:
        rows,boundaries=parse_aux(data,config,start,end)
    except ValueError as exc:
        st.error(str(exc))
        return
    st.caption(f'{boundaries} inícios de tramo/marcadores identificados. Cada vão positivo gera as fases A, B e C, sem conectar os diferentes blocos.')
    mapping={'poste':8,'linear':11,'distance_position':'linha_para'}
    preview_and_import(project,content,filename,key+'_aux',config,rows,mapping,start,end,'aux_tramo',False,0)


def preview_and_import(project,content,filename,key,config,rows,mapping,start,end,phase_mode,percentage,skipped):
    st.write(f'**Prévia: {len(rows)} registros de fase para {config["parque"]}, {config["circuito"]}, nível {config["nivel"]}.**')
    if skipped:
        st.caption(f'{skipped} linhas de outros circuitos/níveis/tipos ou cabeçalhos foram ignoradas.')
    st.dataframe(pd.DataFrame(rows)[['circuito','nivel','fase','condutor','de','para','linear','folga','reserva','corte_fim','origem']].head(300), hide_index=True, width='stretch')
    if len(rows)>300:
        st.caption('A prévia exibe os primeiros 300 registros; todos os registros válidos serão importados.')
    existing = sum(scope_key(r)==scope_key(config) for r in project['trechos'])
    confirmed = True
    if existing:
        confirmed = st.checkbox(f'Substituir os {existing} registros deste circuito, nível, tipo e rota', key=key+'_replace_'+str(scope_key(config)))
    st.caption('Somente este lançamento será atualizado. O estoque e os demais subparques serão preservados. As bobinas serão escolhidas na otimização.')
    st.caption('Importe todos os circuitos, níveis e rotas que serão planejados neste parque. O resumo deste parque passa a ser calculado pelos trechos cadastrados.')
    if st.button('Importar referência para este lançamento', type='primary', disabled=not confirmed, key=key+'_apply'):
        try:
            candidate = apply_reference(project, config, rows, content, filename, mapping, (start,end), phase_mode, percentage)
            persist(project, candidate, 'Referência salva. Confira os trechos e os critérios antes de otimizar.')
        except (ValueError, OSError, sqlite3.Error) as exc:
            st.error(str(exc))


def render_optimization(project, park):
    st.subheader('Otimizar plano de corte')
    st.caption('O cálculo escolhe bobinas e limites de lançamentos entre as estruturas cadastradas, respeitando estoque, folgas, reservas e cortes permitidos. O resultado depende dos critérios e do tempo disponível.')
    key = f"optimize_{project['id']}_{park}_{st.session_state.edit_version}"
    a,b = st.columns(2)
    scope = a.selectbox('Escopo da otimização', ['Este subparque, preservando os demais','Todos os subparques'], key=key+'_scope')
    seconds = b.number_input('Tempo de cálculo [s]', min_value=5,max_value=300,value=30,step=5,key=key+'_time')
    st.caption('Se outros parques ainda não têm bobinas definidas, escolha Todos os subparques para distribuir o estoque em conjunto. Confira as travessias em Critérios de corte.')
    reviewed = st.checkbox('Conferi os pontos de corte, travessias, folgas e reservas dos trechos que serão otimizados.', key=key+'_review')
    if st.button('Calcular bobinas e pontos de corte', type='primary', disabled=not reviewed or not project['trechos'], key=key+'_run'):
        try:
            with st.spinner('Buscando a alocação e verificando o estoque compartilhado…'):
                plan = optimize_scope(project, None if scope=='Todos os subparques' else park, int(seconds))
            all_reviewed = scope=='Todos os subparques' or all(r['parque']==park for r in project['trechos'])
            candidate = dict(project, plano=plan, criterios_confirmados=project.get('criterios_confirmados',False) or all_reviewed)
            persist(project, candidate, plan['status'])
        except (ValueError, OSError, sqlite3.Error) as exc:
            st.error(str(exc))
    plan = project.get('plano')
    if plan and plan.get('fingerprint') == engine.fingerprint(project):
        st.success(plan['status'])
        cuts = [c for c in plan['cortes'] if c['parque']==park]
        st.dataframe(pd.DataFrame(cuts).drop(columns='trechos',errors='ignore'), hide_index=True,width='stretch')
        st.caption('As tabelas e o resumo já exibem as bobinas calculadas. Os arquivos abaixo contêm o plano validado de todos os subparques do projeto.')
        if cuts:
            st.download_button('Baixar Excel final deste subparque',export_subpark(project,park),
                f'Plano_de_Corte_{park.replace("/","-")}.xlsx',
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key=key+'_subpark_final')
        else:
            st.info('Este subparque ainda não tem lançamentos calculados para exportar.')
        a,b=st.columns(2)
        a.download_button('Baixar controle Excel do plano', export(project,'controle'),'Controle_RMT.xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key=key+'_control')
        b.download_button('Baixar entregável Excel do plano', export(project,'entregavel'),'Plano_de_Corte_RMT.xlsx',
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',key=key+'_deliverable')
