"""Converte a sequência I/L do AUX TRAMO em trechos explícitos por fase."""
import math
import io
import colorsys
import re
import xml.etree.ElementTree as ET
import openpyxl
from openpyxl.styles.colors import COLOR_INDEX
from .importer import text, norm, number
from .reference import parse_reference


def is_aux(headers):
    return len(headers)>=12 and norm(headers[8])=='POSTE' and norm(headers[11])=='DISTANCIA TRAMO'


def aux_bounds(data):
    header = next((i for i,row in enumerate(data,1) if is_aux(row)), None)
    if header is None:
        raise ValueError('Cabeçalho POSTE em I e DISTÂNCIA TRAMO em L não encontrado.')
    used = [i for i,row in enumerate(data,1) if i>header and len(row)>8 and text(row[8])]
    if not used:
        raise ValueError('Nenhum poste na coluna I.')
    return header,used[0],used[-1]


def restricted_rows(content, sheet):
    """Reconhece azul em fonte/preenchimento I:L e esforço 1000 em gaveta."""
    w=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    try:
        ws=w[sheet]
        if ws.max_row>30000 or ws.max_column>150:
            raise ValueError('A aba excede o limite de leitura.')
        themes=[]
        if w.loaded_theme:
            root=ET.fromstring(w.loaded_theme)
            scheme=root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
            if scheme is not None:
                themes=[next(iter(e)).get('val') or next(iter(e)).get('lastClr') for e in scheme]
        def blue(color):
            if color is None:return False
            rgb=None
            if color.type=='rgb':rgb=color.rgb
            elif color.type=='indexed' and color.indexed<len(COLOR_INDEX):rgb=COLOR_INDEX[color.indexed]
            elif color.type=='theme' and color.theme<len(themes):rgb=themes[color.theme]
            if not isinstance(rgb,str) or not re.fullmatch('[0-9a-fA-F]{6,8}',rgb):return False
            r,g,b=(int(rgb[-6:][i:i+2],16)/255 for i in (0,2,4))
            if color.tint:
                h,l,s=colorsys.rgb_to_hls(r,g,b)
                l=l*(1+color.tint) if color.tint<0 else l*(1-color.tint)+color.tint
                r,g,b=colorsys.hls_to_rgb(h,l,s)
            h,s,v=colorsys.rgb_to_hsv(r,g,b)
            return .50<=h<=.72 and s>=.20 and v>=.25
        result={}
        for cells in ws.iter_rows(min_col=9,max_col=12):
            if not text(cells[0].value) or norm(cells[0].value)=='POSTE':continue
            marked=any(blue(c.font.color) or (c.fill.patternType=='solid' and blue(c.fill.fgColor)) for c in cells)
            effort=text(cells[1].value).replace(' ','')
            gaveta=norm(cells[2].value)=='GAVETA' and re.search(r'(?:^|/)1000(?:\.0)?$',effort)
            if marked or gaveta:
                result[cells[0].row]='Marcação azul' if marked else 'Esforço 1000 em gaveta'
        return result
    finally:w.close()


def parse_aux(data, config, start, end, restrictions=None):
    """L da linha Para mede o vão entre I da linha anterior e I da atual."""
    if start<1 or end<start or end>len(data):
        raise ValueError('Intervalo inválido.')
    previous=None
    block=0
    physical=[]
    errors=[]
    boundaries=0
    restrictions=restrictions or {}
    initial=text(config.get('origem_inicial')) or 'SE'
    initial_line=int(config.get('linha_inicial',start))
    if not start<=initial_line<=end:raise ValueError('O poste inicial deve estar no intervalo importado.')
    initial_distance=number(config.get('distancia_inicial',0))
    if initial_distance is None or not math.isfinite(initial_distance) or initial_distance<0:
        raise ValueError('Distância inicial inválida.')
    if initial_distance and not initial:
        raise ValueError('Informe a estrutura de origem do vão inicial.')
    for line in range(start,end+1):
        raw=data[line-1]
        pole=text(raw[8]) if len(raw)>8 else ''
        distance=raw[11] if len(raw)>11 else None
        if not pole:
            previous=None
            continue
        if norm(pole)=='POSTE':
            previous=None
            continue
        value=number(distance)
        if value is None or value<0 or not math.isfinite(value):
            errors.append(f'Linha {line}: distância L inválida ou fórmula sem resultado salvo.')
            previous=None
            continue
        placeholder=pole.upper() in ('P.','P','-','—') or pole.startswith('#')
        if placeholder:
            if value!=0:
                errors.append(f'Linha {line}: poste incompleto com distância positiva.')
            previous=None
            boundaries+=1
            continue
        if value==0:
            block+=1
            boundaries+=1
            # L=0 marks a nonexistent pole: the next positive span starts at SE.
            previous=(initial,0)
            continue
        if previous is None:
            if line==start and line==initial_line and initial:
                previous=(initial,0)
                block+=1
            else:
                errors.append(f'Linha {line}: distância positiva sem poste anterior no intervalo. Inclua a linha de início do tramo ou informe a origem inicial (SE).')
                continue
        if previous[0]==pole:
            errors.append(f'Linha {line}: mesmo poste nas duas extremidades.')
        else:
            physical.append((previous[0],pole,value,previous[1],line,block))
        previous=(pole,line)
    if errors:
        raise ValueError('\n'.join(errors[:30]))
    if not physical:
        raise ValueError('Nenhum vão positivo entre postes foi encontrado no intervalo.')
    result=[]
    offset=0
    # Reusa a validação de números, parâmetros e geração A/B/C da importação normal.
    for block_id in sorted({r[5] for r in physical}):
        records=[r for r in physical if r[5]==block_id]
        route=f"{config['rota']} · TRAMO {block_id}"
        rows,_=parse_reference([[r[0],r[1],r[2]] for r in records],dict(de=0,para=1,linear=2),
            {**config,'rota':route},1,len(records),'expand')
        for i, record in enumerate(records):
            for row in rows[i*3:i*3+3]:
                row['origem']=f"{config['sheet']}!I{record[3]}→I{record[4]} / L{record[4]}"
                row['referencia_de_linha']=record[3]
                row['referencia_para_linha']=record[4]
                row['referencia_distancia_linha']=record[4]
                row['sem_corte_de']=bool(restrictions.get(record[3]))
                row['sem_corte_para']=bool(restrictions.get(record[4]))
                row['restricao_estrutura']='; '.join(dict.fromkeys(str(restrictions[i]) for i in record[3:5] if restrictions.get(i)))
                if row['sem_corte_para']:row['corte_fim']='PROIBIDO'
                if record[3]==0:
                    row['origem']=f"{config['sheet']}!I{record[4]} / origem inicial informada"
                if row['restricao_estrutura']:
                    row['observacao']='Sem corte na estrutura: '+row['restricao_estrutura']
                # O escopo continua sendo a rota escolhida; separação física via identificador de bloco.
                row['rota']=config['rota']
                row['bloco_tramo']=str(block_id)
                row['ordem']=offset+i+1
        result.extend(rows)
        offset+=len(records)
    return result,boundaries
