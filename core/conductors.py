"""Catálogo de condutores por obra e aplicação em trechos selecionados."""
import copy
import io
import json
from pathlib import Path
import openpyxl
from .importer import norm, text
from . import engine


def parse_catalog(content, filename='Dados Condutores.xlsx'):
    w=openpyxl.load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    result=[]
    try:
        for ws in w:
            if ws.max_row>30000 or ws.max_column>150:
                raise ValueError('A aba excede o limite de leitura.')
            kind=None
            headers=None
            for line,raw in enumerate(ws.values,1):
                row=list(raw)
                label=norm(row[1]) if len(row)>1 else ''
                if label in ('CABOS SUBTERRANEOS','CABOS AEREOS'):
                    kind='SUBTERRÂNEO' if label=='CABOS SUBTERRANEOS' else 'AÉREO'
                    headers=None
                elif label=='DESCRICAO' and kind:
                    headers=[text(v).replace('\n',' ') for v in row[1:10]]
                elif kind and headers:
                    description=text(row[1]) if len(row)>1 else ''
                    if not description or label=='N/A':
                        kind=None
                        headers=None
                        continue
                    if not isinstance(row[1],str):
                        raise ValueError(f'{ws.title}!B{line}: descrição inválida.')
                    properties={h:v for h,v in zip(headers[1:],row[2:10]) if h and h!='Coluna1'}
                    result.append(dict(descricao=description,tipo=kind,dados=properties,
                        origem=f'{filename} · {ws.title}!B{line}'))
        if not result:raise ValueError('Não foram encontrados os blocos CABOS SUBTERRÂNEOS / CABOS AÉREOS com DESCRIÇÃO na coluna B.')
        if len({(r['tipo'],r['descricao']) for r in result})!=len(result):
            raise ValueError('O arquivo tem descrições de condutores duplicadas no mesmo tipo.')
        return result
    finally:w.close()


def catalog(project):
    source=json.loads(Path(__file__).with_name('condutores_padrao.json').read_text(encoding='utf-8'))
    records={(r['tipo'],r['descricao']):r for r in source}
    records.update({(r['tipo'],r['descricao']):r for r in project.get('condutores',[])})
    return sorted(records.values(),key=lambda r:(r['tipo'],r['descricao']))


def options(project, installation):
    return sorted({r['descricao'] for r in catalog(project) if r['tipo']==installation} |
        {r['condutor'] for r in project['bobinas']+project['trechos'] if r.get('tipo')==installation and r.get('condutor')})


def register(project, records):
    candidate=copy.deepcopy(project)
    merged={(r['tipo'],r['descricao']):r for r in candidate.get('condutores',[])}
    for r in records:
        description=text(r.get('descricao'))
        if not description or r.get('tipo') not in ('AÉREO','SUBTERRÂNEO'):
            raise ValueError('Informe descrição e instalação válidas.')
        merged[(r['tipo'],description)]={**r,'descricao':description}
    candidate['condutores']=list(merged.values())
    return candidate


def assign(project, park, ids, description, installation):
    if not ids:raise ValueError('Selecione ao menos um trecho.')
    if description not in options(project,installation):raise ValueError('Selecione um condutor cadastrado.')
    chosen=set(ids)
    existing={r['id']:r for r in project['trechos']}
    if any(i not in existing or existing[i]['parque']!=park or existing[i]['tipo']!=installation for i in chosen):
        raise ValueError('Trechos inválidos ou pertencentes a outro subparque.')
    candidate=copy.deepcopy(project)
    plan=project.get('plano')
    assigned={i:c['bobina'] for c in plan['cortes'] for i in c['trechos']} if plan and plan.get('fingerprint')==engine.fingerprint(project) else {}
    reels={r['id']:r for r in project['bobinas']}
    for row in candidate['trechos']:
        if row['id'] in assigned:row['bobina_original']=assigned[row['id']]
        if row['id'] not in chosen:continue
        row['condutor']=description
        for field in ('bobina_original','fixa'):
            reel=reels.get(row.get(field))
            if (reel and (reel['condutor']!=description or reel['tipo']!=installation)) or (row.get(field) and reel is None):
                row[field]=''
    # Uma mudança de condutor precisa de um ponto físico de corte permitido.
    blocked=engine.blocked_nodes(candidate['trechos'])
    for chain in engine.segments(candidate['trechos']):
        if not any(r['id'] in chosen for r in chain):continue
        if not engine.start_allowed(chain[0],blocked) or not engine.end_allowed(chain[-1],blocked):
            raise ValueError('A seleção termina ou começa em uma estrutura sem corte. Inclua os trechos adjacentes até um ponto permitido.')
    candidate['plano']=None
    candidate['criterios_confirmados']=False
    return candidate
