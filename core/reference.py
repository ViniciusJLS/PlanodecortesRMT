"""Referência de traçado por subparque; não importa estoque nem executa fórmulas."""
import base64
import copy
import hashlib
import io
import uuid
from collections import defaultdict
from datetime import datetime, timezone
from decimal import Decimal
import openpyxl
from . import engine
from .importer import norm, number, text
from .inventory import park_id
from .subparks import CIRCUITS, LEVELS, circuit_id

FIELDS = {
    'de':'De', 'para':'Para', 'linear':'Distância linear [m]', 'fase':'Fase',
    'condutor':'Condutor', 'circuito':'Circuito', 'nivel':'Nível', 'tipo':'Tipo',
    'folga':'Folga', 'reserva':'Reserva total [m]', 'sobra_caixa':'Sobra total caixas [m]',
    'sobra_poste':'Sobra total postes [m]', 'sobra_turbina':'Sobra saída turbina [m]',
    'ordem':'Ordem', 'corte_fim':'Corte no fim', 'observacao':'Observações',
}


def sheets(content):
    with io.BytesIO(content) as stream:
        w = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            return w.sheetnames
        finally:
            w.close()


def read_sheet(content, sheet):
    w = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = w[sheet]
        if ws.max_row > 30000 or ws.max_column > 150:
            raise ValueError('A aba excede 30 mil linhas ou 150 colunas. Exporte apenas a tabela de traçado para um novo Excel.')
        return list(ws.values)
    finally:
        w.close()


def guess_header(rows):
    for i, row in enumerate(rows[:50], 1):
        labels = {norm(v) for v in row}
        if 'DE' in labels and 'PARA' in labels:
            return i
    return 1


def guess_mapping(headers, level):
    aliases = {
        'de':['DE','INICIO','ESTRUTURA INICIAL'], 'para':['PARA','FIM','ESTRUTURA FINAL'],
        'linear':['DISTANCIA LINEAR','DISTANCIA LINEAR [M]','LINEAR [M]','LINEAR'],
        'fase':['FASE'], 'condutor':[f'DESCRICAO (NIVEL {level})','CONDUTOR','DESCRICAO'],
        'circuito':[f'CIRCUITO (NIVEL {level})','CIRCUITO'], 'nivel':['NIVEL'], 'tipo':['TIPO'],
        'folga':['FOLGA DESNIVEL','FOLGA','FOLGA (FRACAO)','FOLGA [%]'],
        'reserva':['RESERVA [M]','RESERVA','RESERVA TOTAL [M]'],
        'sobra_caixa':['SOBRA CAIXA'], 'sobra_poste':['SOBRA POSTE'],
        'sobra_turbina':['SOBRA SAIDA TURBINA'], 'ordem':['ORDEM'],
        'corte_fim':['CORTE NO FIM','CORTE_FIM'], 'observacao':['OBSERVACOES','OBSERVACAO'],
    }
    labels = [norm(v) for v in headers]
    return {k:next((labels.index(a) for a in aliases[k] if a in labels), None) for k in FIELDS}


def scope_key(row):
    return (park_id(row['parque']), circuit_id(row['circuito']), str(row['nivel']), row['tipo'], row['rota'])


def parse_reference(rows, mapping, config, start, end, phase_mode='expand', percent=False):
    """Validação completa antes de alterar o projeto; linhas informadas são 1-based."""
    if config['circuito'] not in CIRCUITS or str(config['nivel']) not in LEVELS:
        raise ValueError('Escolha circuito C01–C34 e nível 1–3.')
    if not text(config.get('parque')) or not text(config.get('rota')):
        raise ValueError('Informe subparque e rota.')
    if config['tipo'] not in ('AÉREO','SUBTERRÂNEO'):
        raise ValueError('Tipo de instalação inválido.')
    if any(mapping.get(k) is None for k in ('de','para','linear')):
        raise ValueError('Associe as colunas De, Para e Distância linear.')
    used = [v for v in mapping.values() if v is not None]
    if len(used) != len(set(used)):
        raise ValueError('Uma coluna foi associada a mais de um campo. Revise o mapeamento.')
    if start < 1 or end < start or end > len(rows):
        raise ValueError('Intervalo de linhas inválido.')
    if phase_mode not in ('expand','column') or (phase_mode == 'column' and mapping.get('fase') is None):
        raise ValueError('Selecione a coluna de fase ou gere A/B/C a partir de uma linha por trecho.')
    counters = defaultdict(int)
    result, errors, skipped = [], [], 0
    for line in range(start, end+1):
        raw = rows[line-1]
        def get(field):
            col = mapping.get(field)
            return raw[col] if col is not None and col < len(raw) else None
        if all(v in (None,'') for v in raw):
            continue
        if norm(get('de')) == 'DE' and norm(get('para')) == 'PARA':
            skipped += 1
            continue
        try:
            source_circuit = circuit_id(get('circuito'))
            source_level = text(get('nivel'))
            source_type = norm(get('tipo'))
            if ((source_circuit and source_circuit != config['circuito']) or
                (source_level and source_level != str(config['nivel'])) or
                (source_type and source_type != norm(config['tipo']))):
                skipped += 1
                continue
            if mapping.get('circuito') is not None and source_circuit not in CIRCUITS:
                raise ValueError('circuito vazio ou inválido na coluna associada')
            if phase_mode == 'expand' and text(get('fase')):
                raise ValueError('a referência já contém fases; escolha o modo Uma linha por fase')
            phases = ['A','B','C'] if phase_mode == 'expand' else [text(get('fase')).upper()]
            if any(p not in ('A','B','C') for p in phases):
                raise ValueError('fase deve ser A, B ou C')
            a, b = text(get('de')), text(get('para'))
            if not a or not b or a.startswith('#') or b.startswith('#') or a == b:
                raise ValueError('De/Para vazios, iguais ou com erro Excel')
            conductor = text(get('condutor')) or text(config.get('condutor'))
            if not conductor or conductor.startswith('#'):
                raise ValueError('informe o condutor no Excel ou na configuração')
            def numeric(field, fallback=0):
                value = get(field)
                if value in (None,''):
                    if mapping.get(field) is not None:
                        raise ValueError(f'{FIELDS[field]} vazio; se for fórmula, recalcule e salve o Excel')
                    value = fallback
                n = number(value)
                if n is None or not engine.dec(n).is_finite() or n < 0:
                    raise ValueError(f'{FIELDS[field]} inválido')
                return n
            linear = numeric('linear')
            fraction = numeric('folga', config['folga'])
            if percent and mapping.get('folga') is not None:
                fraction /= 100
            if linear <= 0 or fraction > 1:
                raise ValueError('distância deve ser positiva e folga entre 0 e 100%')
            details = {f:numeric(f) for f in ('sobra_caixa','sobra_poste','sobra_turbina')}
            if mapping.get('reserva') is not None and any(mapping.get(f) is not None for f in details):
                raise ValueError('mapeie a reserva total OU as sobras detalhadas, para não duplicar reservas')
            reserve = numeric('reserva', config.get('reserva',0)) if not any(mapping.get(f) is not None for f in details) else sum(details.values())
            cut = norm(get('corte_fim')) or config.get('corte_fim','PERMITIDO')
            if cut not in ('PERMITIDO','PROIBIDO','OBRIGATORIO'):
                raise ValueError('corte no fim deve ser PERMITIDO, PROIBIDO ou OBRIGATORIO')
            for phase in phases:
                counters[phase] += 1
                order = numeric('ordem', counters[phase])
                if order <= 0:
                    raise ValueError('ordem deve ser positiva')
                result.append(dict(id=str(uuid.uuid4()), parque=config['parque'], circuito=config['circuito'],
                    nivel=str(config['nivel']), tipo=config['tipo'], rota=config['rota'], fase=phase,
                    condutor=conductor, ordem=order, de=a, para=b, linear=linear, folga=fraction,
                    reserva=reserve, corte_fim=cut, fixa='', bobina_original='', observacao=text(get('observacao')),
                    origem=f"{config.get('sheet','Referência')}!{line}", **details,
                    reserva_outros=reserve-sum(details.values())))
        except (ValueError, TypeError, ArithmeticError) as exc:
            errors.append(f'Linha {line}: {exc}.')
    if errors:
        raise ValueError('\n'.join(errors[:30]))
    if not result:
        raise ValueError('Nenhum trecho corresponde ao circuito, nível, tipo e intervalo selecionados.')
    seen = set()
    triples = defaultdict(set)
    for r in result:
        key = (r['fase'],r['ordem'])
        if key in seen:
            raise ValueError('Há ordens repetidas para a mesma fase. Revise a referência.')
        seen.add(key)
        triples[(r['ordem'],r['de'],r['para'])].add(r['fase'])
    if any(phases != {'A','B','C'} for phases in triples.values()):
        raise ValueError('Cada trecho precisa das fases A, B e C com a mesma ordem e De/Para. Confira linhas ausentes ou o intervalo selecionado.')
    return result, skipped


def register_park(project, name):
    name = text(name)
    if not name:
        raise ValueError('Informe o nome do subparque.')
    existing = {park_id(r['parque']):r['parque'] for r in project['trechos']}
    existing.update({park_id(n):n for n in project.get('subparques',[])})
    name = existing.get(park_id(name), name)
    candidate = copy.deepcopy(project)
    candidate['subparques'] = sorted(set(candidate.get('subparques',[])) | {name})
    return candidate, name


def apply_reference(project, config, rows, content, filename, mapping, interval, phase_mode, percent):
    candidate, name = register_park(project, config['parque'])
    config = {**config, 'parque':name}
    target = scope_key(config)
    if not rows or any(scope_key(r) != target for r in rows):
        raise ValueError('A prévia não corresponde à configuração selecionada.')
    # Ao importar um parque antes só presente no resumo, seu histórico vira escopo de planejamento.
    for reel in candidate['bobinas']:
        scope = {park_id(k) for k in reel.get('parques_replanejados',[])}
        if park_id(name) not in scope:
            imported = {park_id(k):v for k,v in reel.get('consumo_importado_parques',{}).items()}
            old = imported.get(park_id(name), 0)
            if old is None or engine.dec(old) < 0:
                raise ValueError(f"{reel['id']}: consumo histórico deste parque pendente; corrija antes de importar.")
            previous = engine.dec(reel.get('utilizado',0)) - engine.dec(old)
            if previous < 0:
                raise ValueError(f"{reel['id']}: histórico incompatível com consumo anterior; revise Bobinas.")
            reel['utilizado'] = float(previous)
            reel['parques_replanejados'] = sorted(scope | {park_id(name)})
    plan = project.get('plano')
    if plan and plan.get('fingerprint') == engine.fingerprint(project):
        assigned = {ident:c['bobina'] for c in plan['cortes'] for ident in c['trechos']}
        for row in candidate['trechos']:
            if row['id'] in assigned:
                row['bobina_original'] = assigned[row['id']]
    candidate['trechos'] = [r for r in candidate['trechos'] if scope_key(r) != target] + copy.deepcopy(rows)
    for row in candidate['trechos']:
        row['circuito'] = circuit_id(row['circuito'])
    digest = hashlib.sha256(content).hexdigest()
    refs = candidate.setdefault('referencias_subparques', {})
    refs.setdefault(name, {})[' | '.join(target[1:])] = dict(
        filename=filename, sha256=digest, sheet=config['sheet'], config=config, mapping=mapping,
        interval=list(interval), phase_mode=phase_mode, percent=percent,
        uploaded=datetime.now(timezone.utc).isoformat(), registros=len(rows))
    candidate.setdefault('arquivos_referencia', {})[digest] = base64.b64encode(content).decode('ascii')
    used = {v['sha256'] for park in refs.values() for v in park.values()}
    candidate['arquivos_referencia'] = {k:v for k,v in candidate['arquivos_referencia'].items() if k in used}
    candidate['plano'] = None
    candidate['criterios_confirmados'] = False
    return candidate


def optimize_scope(project, park=None, timeout=30):
    """Otimiza o parque usando o saldo após reservar os lançamentos dos demais."""
    if park is None:
        return engine.optimize(project, timeout)
    selected = [r for r in project['trechos'] if r['parque']==park]
    if not selected:
        raise ValueError('Importe os trechos deste subparque antes de otimizar.')
    other = [r for r in project['trechos'] if r['parque']!=park]
    fixed = []
    if other:
        plan = project.get('plano')
        if plan and plan.get('fingerprint') == engine.fingerprint(project):
            fixed = copy.deepcopy([c for c in plan['cortes'] if c['parque']!=park])
        else:
            if any(not r.get('bobina_original') for r in other):
                raise ValueError('Há trechos de outros parques sem bobina. Use Todos os subparques para alocar o estoque compartilhado.')
            external = copy.deepcopy(project)
            external['trechos'] = other
            fixed = engine.consolidate_existing(external)['cortes']
    reserved = defaultdict(Decimal)
    for cut in fixed:
        reserved[cut['bobina']] += engine.dec(cut['projeto'])
    local = copy.deepcopy(project)
    local['trechos'] = selected
    for reel in local['bobinas']:
        reel['utilizado'] = float(engine.dec(reel.get('utilizado',0)) + reserved[reel['id']])
    result = engine.optimize(local, timeout)
    cuts = fixed + result['cortes']
    for i, cut in enumerate(cuts,1):
        cut['id'] = f'C{i:04d}'
    errors = engine.validate(project, cuts)
    if errors:
        raise ValueError('\n'.join(errors[:30]))
    result.update(cortes=cuts, fingerprint=engine.fingerprint(project),
        status=f"{result['status']} — {park}, preservando os demais subparques")
    return result


def reference_template():
    from .exporter import workbook
    return workbook([('Traçado','REFERÊNCIA DO SUBPARQUE',
        'Uma linha por trecho físico; A/B/C serão geradas na importação. Informe circuito, nível e condutor no aplicativo.',
        ['De','Para','Distância linear [m]','Corte no fim','Observações'],
        [['P.0/1A','P.0/2A',100.25,'PERMITIDO','Exemplo: substitua pelos dados reais'],
         ['P.0/2A','P.0/3A',125.5,'PERMITIDO','']], [25,25,25,24,60])])
