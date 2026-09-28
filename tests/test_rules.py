import copy
import io
import unittest
from pathlib import Path
import openpyxl
from core import engine, project, importer
from core.exporter import export


class Rules(unittest.TestCase):
    def setUp(self):
        self.p=project.demo()

    def test_actual_length_replaces_tolerance(self):
        reel=dict(id='B',nominal=1500,real=None,utilizado=100)
        self.assertEqual(float(engine.available(reel)),1355)
        reel['real']=1487
        self.assertEqual(float(engine.available(reel)),1387)
        reel['real']=0
        self.assertEqual(float(engine.base(reel)),0)

    def test_round_once_and_decimal_precision(self):
        rows=[dict(linear=10.1,folga=.05,reserva=0),dict(linear=10.1,folga=.05,reserva=0)]
        self.assertEqual(sum(engine.length([r]) for r in rows),22)
        self.assertEqual(engine.length(rows),22)
        rows=[dict(linear=10.01,folga=.05,reserva=0)]*3
        self.assertEqual(sum(engine.length([r]) for r in rows),33)
        self.assertEqual(engine.length(rows),32)
        self.assertEqual(engine.length([dict(linear=20,folga=.05,reserva=0)]),21)

    def test_complete_optimized_solution_and_xlsx(self):
        self.p['plano']=engine.optimize(self.p,8)
        cuts=self.p['plano']['cortes']
        self.assertEqual(engine.validate(self.p,cuts),[])
        for c in cuts:
            self.assertNotEqual(c['para'],'E03')
        for kind in ['controle','entregavel']:
            content=export(self.p,kind)
            w=openpyxl.load_workbook(io.BytesIO(content),data_only=True)
            self.assertIn('Resumo',w.sheetnames)
            reported=sum(float(w['Resumo'].cell(i,3).value or 0) for i in range(5,w['Resumo'].max_row+1))
            self.assertEqual(reported,sum(c['projeto'] for c in cuts))
            for sheet in w:
                for row in sheet:
                    self.assertFalse(any(c.data_type=='e' for c in row))
            w.close()
        self.p['trechos'][0]['linear']+=1
        with self.assertRaises(ValueError):export(self.p,'controle')

    def test_auditor_rejects_duplicate_and_overflow(self):
        self.p['plano']=engine.optimize(self.p,8)
        cuts=self.p['plano']['cortes']
        self.assertTrue(engine.validate(self.p,cuts+cuts[:1]))
        self.p['bobinas'][0]['real']=1
        for c in cuts:c['bobina']=self.p['bobinas'][0]['id']
        self.assertTrue(any('excede' in e for e in engine.validate(self.p,cuts)))

    def test_mandatory_cut_fixed_reel_and_infeasible(self):
        self.p['trechos']=self.p['trechos'][:6]
        self.p['trechos'][2]['corte_fim']='OBRIGATORIO'
        self.p['trechos'][0]['fixa']='DEMO-MAG-002'
        result=engine.optimize(self.p,5)
        self.assertTrue(any(c['para']=='E04' for c in result['cortes']))
        cut=next(c for c in result['cortes'] if 'T01-A' in c['trechos'])
        self.assertEqual(cut['bobina'],'DEMO-MAG-002')
        for b in self.p['bobinas']:b['real']=1
        with self.assertRaises(ValueError):engine.optimize(self.p,2)

    def test_disconnected_same_reel_not_merged(self):
        a=copy.deepcopy(self.p['trechos'][0]);b=copy.deepcopy(self.p['trechos'][1])
        a.update(bobina_original='DEMO-MAG-001',corte_fim='PERMITIDO')
        b.update(de='OUTRO',para='FINAL',bobina_original='DEMO-MAG-001',corte_fim='PERMITIDO')
        self.p['trechos']=[a,b]
        self.assertEqual(len(engine.consolidate_existing(self.p)['cortes']),2)

    def test_original_import_swapped_columns_and_external_usage(self):
        path=Path('referencias/controle_original.xlsx')
        if not path.exists():self.skipTest('Referência não incluída')
        rows,reels,warnings=importer.import_control(path.read_bytes(),['RSA-12'])
        a=next(r for r in rows if r['id']=='RSA-12-4-N2')
        self.assertEqual(a['circuito'],'32')
        self.assertEqual(a['bobina_original'],'EDIV-MAG-103')
        self.assertEqual(a['linear'],26.5)
        reel=next(r for r in reels if r['id']=='EDIV-095-007')
        self.assertEqual(reel['utilizado'],0)
        self.assertEqual(float(engine.base(reel)),1843)
        self.assertIsNone(reel['real'])


if __name__=='__main__':unittest.main()
